"""خروجی Excel و PDF از «جدول استاندارد» گزارش‌ها (ر.ک. app/services/reports.py).

- Excel: openpyxl، راست‌به‌چپ، سلول‌های متنی با پیشوند خطرناک (= + - @) خنثی می‌شوند (CSV/formula injection).
- PDF: reportlab، A4 افقی، فارسی با arabic-reshaper + python-bidi و فونت TTF (PDF_FONT_PATH یا Vazirmatn/DejaVu).
  اگر کتابخانه‌ی شکل‌دهی فارسی نصب نباشد خطای روشن می‌دهد؛ متن خراب تولید نمی‌شود.
"""
from __future__ import annotations
import io
import os
from datetime import date,datetime
from decimal import Decimal
from app.core.config import ROOT
from app.services.jalali import fmt_jalali

FA=str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹')
def _fmt(v,kind,fa=False):
 """مقدار نمایشی (برای PDF و Excel متنی)"""
 if v is None or v=='': return '—' if kind!='text' else ''
 if kind in ('date','datetime'):
  try:
   if isinstance(v,(date,datetime)): s=fmt_jalali(v)
   else: s=fmt_jalali(date.fromisoformat(str(v)[:10]))
  except ValueError: s=str(v)
 elif kind in ('money','int'): s=f'{Decimal(str(v)):,.0f}'
 elif kind in ('qty','pct'):
  x=Decimal(str(v)).quantize(Decimal('0.01'));s=f'{x:,.2f}'.rstrip('0').rstrip('.') if x!=x.to_integral() else f'{x:,.0f}'
  if kind=='pct': s+='٪' if fa else '%'
 else: s=str(v)
 return s.translate(FA).replace(',','٬').replace('.','٫') if fa else s
def _cell(v,kind):
 if v is None: return None
 if kind in ('money','int','qty','pct') and not isinstance(v,str): return float(v)
 if kind in ('date','datetime'): return _fmt(v,kind)
 if isinstance(v,Decimal): return float(v)
 if isinstance(v,(date,datetime)): return v.isoformat()
 if isinstance(v,str) and v[:1] in ('=','+','-','@'): return "'"+v
 return v
def to_xlsx(t:dict)->bytes:
 from openpyxl import Workbook
 from openpyxl.styles import Font
 wb=Workbook();ws=wb.active;ws.title=(t['title'] or 'report')[:31].replace('/','-');ws.sheet_view.rightToLeft=True
 cols=t['columns'];ws.append([_cell(t['title'],'text')]);ws['A1'].font=Font(bold=True,size=13)
 if t.get('filters'): ws.append([_cell('فیلترها: '+'، '.join(f'{k}={v}' for k,v in t['filters'].items()),'text')])
 ws.append([]);hr=ws.max_row+1;ws.append([c['label'] for c in cols])
 for c in ws[hr]: c.font=Font(bold=True)
 for r in t['rows']: ws.append([_cell(r.get(c['key']),c['kind']) for c in cols])
 if t.get('summary'):
  ws.append([])
  for lab,val,kind in t['summary']: ws.append([_cell(lab,'text'),_cell(val,kind)])
 for i,c in enumerate(cols,1): ws.column_dimensions[ws.cell(row=hr,column=i).column_letter].width=max(12,min(40,len(c['label'])+4))
 b=io.BytesIO();wb.save(b);return b.getvalue()
def _font_path():
 for p in [os.environ.get('PDF_FONT_PATH',''),os.path.join(ROOT,'web','fonts','Vazirmatn-Regular.ttf'),'/usr/share/fonts/truetype/vazirmatn/Vazirmatn-Regular.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf','/usr/share/fonts/dejavu/DejaVuSans.ttf']:
  if p and os.path.isfile(p): return p
 raise RuntimeError('فونت فارسی برای PDF پیدا نشد؛ مسیر را در PDF_FONT_PATH بدهید یا بسته‌ی fonts-dejavu-core را نصب کنید')
def _shaper():
 try:
  import arabic_reshaper
  from bidi.algorithm import get_display
 except ImportError as e: raise RuntimeError('برای PDF فارسی بسته‌های arabic-reshaper و python-bidi لازم است') from e
 rs=arabic_reshaper.ArabicReshaper({'delete_harakat':False,'support_ligatures':False})
 return lambda s:get_display(rs.reshape(str(s)))
def to_pdf(t:dict,company:str='')->bytes:
 from reportlab.lib import colors
 from reportlab.lib.pagesizes import A4,landscape
 from reportlab.lib.styles import ParagraphStyle
 from reportlab.lib.units import mm
 from reportlab.pdfbase import pdfmetrics
 from reportlab.pdfbase.ttfonts import TTFont
 from reportlab.platypus import Paragraph,SimpleDocTemplate,Spacer,Table,TableStyle
 sh=_shaper();fp=_font_path()
 if 'RptFont' not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont('RptFont',fp))
 base=ParagraphStyle('b',fontName='RptFont',fontSize=8,leading=11,alignment=2);ttl=ParagraphStyle('t',parent=base,fontSize=14,leading=19);sm=ParagraphStyle('s',parent=base,fontSize=9,leading=13)
 cols=list(reversed(t['columns']))   # راست‌به‌چپ: ستون اول سمت راست
 P=lambda s,st=base:Paragraph(sh(s).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;') if s!='' else '',st)
 head=[P(c['label']) for c in cols];body=[[P(_fmt(r.get(c['key']),c['kind'],True)) for c in cols] for r in t['rows']]
 W=landscape(A4)[0]-20*mm;cw=W/max(len(cols),1)
 tb=Table([head]+body,colWidths=[cw]*len(cols),repeatRows=1)
 tb.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8eefc')),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#b8c2d8')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f6f8fd')])]))
 story=[P(t['title'],ttl)]
 meta=[x for x in [company,'تاریخ گزارش: '+_fmt(date.today(),'date',True)] if x]
 if t.get('filters'): meta.append('فیلترها: '+'، '.join(f'{k}={v}' for k,v in t['filters'].items()))
 story+=[P(' — '.join(meta),sm),Spacer(1,4*mm),tb,Spacer(1,4*mm)]
 for lab,val,kind in t.get('summary') or []: story.append(P(lab+(': '+_fmt(val,kind,True) if val!='' else ''),sm))
 b=io.BytesIO();SimpleDocTemplate(b,pagesize=landscape(A4),leftMargin=10*mm,rightMargin=10*mm,topMargin=10*mm,bottomMargin=10*mm,title=t['title']).build(story);return b.getvalue()

def invoice_pdf(d:dict)->bytes:
 """PDF فاکتور فروش A4 عمودی از خروجی invoice_detail (راست‌به‌چپ، فارسی، تاریخ شمسی)."""
 from reportlab.lib import colors
 from reportlab.lib.pagesizes import A4
 from reportlab.lib.styles import ParagraphStyle
 from reportlab.lib.units import mm
 from reportlab.pdfbase import pdfmetrics
 from reportlab.pdfbase.ttfonts import TTFont
 from reportlab.platypus import Paragraph,SimpleDocTemplate,Spacer,Table,TableStyle
 sh=_shaper();fp=_font_path()
 if 'RptFont' not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont('RptFont',fp))
 base=ParagraphStyle('ib',fontName='RptFont',fontSize=9,leading=13,alignment=2);ttl=ParagraphStyle('it',parent=base,fontSize=15,leading=21);sm=ParagraphStyle('is',parent=base,fontSize=8,leading=12)
 esc=lambda x:str(x).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')
 P=lambda x,st=base:Paragraph(esc(sh(x)) if x not in ('',None) else '',st)
 inv=d['invoice'];co=d.get('company') or {};c=d.get('customer') or {}
 money=lambda v:_fmt(v,'money',True);qty=lambda v:_fmt(v,'qty',True)
 created=inv.get('created_at');date_s=_fmt(created,'date',True) if created else ''
 story=[P(co.get('name',''),ttl),P(f"فاکتور فروش شماره {inv['number']}".translate(FA)+(f" — تاریخ: {date_s}" if date_s else ''),base),Spacer(1,3*mm)]
 info=[[P('مشتری: '+str(c.get('name',''))+(f" ({c['code']})" if c.get('code') else '')),P('فروشنده: '+str(d.get('seller','')))],[P('نشانی: '+str(c.get('address') or '—')),P('تلفن: '+str(c.get('phone') or '—'))],[P('انبار: '+str(d.get('warehouse') or '—')),P('سفارش: '+str(d.get('order_number','')))]]
 it=Table(info,colWidths=[95*mm,95*mm]);it.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),2)]));story+=[it,Spacer(1,3*mm)]
 heads=['ردیف','شرح کالا','واحد','تعداد','قیمت واحد','تخفیف','مبلغ'];cols=list(reversed(heads))
 rows=[[P(h) for h in cols]]
 for i,l in enumerate(d.get('lines') or [],1):
  cells=[str(i).translate(FA),l.get('description') or '',l.get('unit') or '',qty(l.get('quantity')),money(l.get('unit_price')),money(l.get('discount') or 0),money(l.get('line_total'))]
  rows.append([P(x) for x in reversed(cells)])
 cw=[28*mm,22*mm,24*mm,16*mm,16*mm,52*mm,12*mm]
 tb=Table(rows,colWidths=cw,repeatRows=1)
 tb.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8eefc')),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#b8c2d8')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f6f8fd')])]))
 story+=[tb,Spacer(1,4*mm)]
 tot=[('جمع کل اقلام',inv.get('subtotal')),('مالیات بر ارزش افزوده',inv.get('tax')),('هزینهٔ حمل',inv.get('shipping')),('مبلغ فاکتور',inv.get('total')),('برگشت از فروش',inv.get('returned')),('پرداخت‌شده',inv.get('paid')),('مانده',d.get('remaining'))]
 tt=Table([[P(money(v or 0)),P(k)] for k,v in tot],colWidths=[40*mm,45*mm],hAlign='RIGHT')
 tt.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.25,colors.HexColor('#b8c2d8')),('BACKGROUND',(0,3),(-1,3),colors.HexColor('#e8eefc'))]));story+=[tt,Spacer(1,4*mm)]
 pays=d.get('payments') or []
 if pays: story.append(P('پرداخت‌ها: '+'، '.join(f"{money(p['amount'])} ({ {'cash':'نقد','card':'کارت','transfer':'حواله','check':'چک'}.get(p.get('method'),p.get('method')) })" for p in pays),sm))
 if inv.get('note'): story.append(P('یادداشت: '+str(inv['note']),sm))
 if inv.get('terms'): story.append(P('شرایط: '+str(inv['terms']),sm))
 b=io.BytesIO();SimpleDocTemplate(b,pagesize=A4,leftMargin=10*mm,rightMargin=10*mm,topMargin=10*mm,bottomMargin=10*mm,title='فاکتور '+str(inv['number'])).build(story);return b.getvalue()

def payslip_pdf(d:dict)->bytes:
 """فیش حقوقی A5 عمودی از خروجی payslip_data (راست‌به‌چپ، فارسی)."""
 from reportlab.lib import colors
 from reportlab.lib.pagesizes import A5
 from reportlab.lib.styles import ParagraphStyle
 from reportlab.lib.units import mm
 from reportlab.pdfbase import pdfmetrics
 from reportlab.pdfbase.ttfonts import TTFont
 from reportlab.platypus import Paragraph,SimpleDocTemplate,Spacer,Table,TableStyle
 sh=_shaper();fp=_font_path()
 if 'RptFont' not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont('RptFont',fp))
 base=ParagraphStyle('pb',fontName='RptFont',fontSize=10,leading=15,alignment=2);ttl=ParagraphStyle('pt',parent=base,fontSize=14,leading=20)
 esc=lambda x:str(x).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')
 P=lambda x,st=base:Paragraph(esc(sh(x)) if x not in ('',None) else '',st)
 money=lambda v:_fmt(v,'money',True);l=d['line'];e=d['employee']
 rows=[('حقوق پایه',l['base']),('پاداش',l['bonus']),(f"کسر مرخصی بدون‌حقوق ({_fmt(l['unpaid_days'],'qty',True)} روز)",l['deduction']),('خالص پرداختی',l['net'])]
 tb=Table([[P(money(v)),P(k)] for k,v in rows],colWidths=[45*mm,65*mm],hAlign='RIGHT')
 tb.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.25,colors.HexColor('#b8c2d8')),('BACKGROUND',(0,3),(-1,3),colors.HexColor('#e8eefc')),('VALIGN',(0,0),(-1,-1),'MIDDLE')]))
 y,m=d['period'].split('-')
 story=[P(d['company'].get('name',''),ttl),P('فیش حقوقی دورهٔ '+f'{y}-{m}'.translate(FA)),Spacer(1,3*mm),P('نام: '+str(e['name'])),P('سمت: '+str(e.get('position') or '—')),Spacer(1,4*mm),tb]
 b=io.BytesIO();SimpleDocTemplate(b,pagesize=A5,leftMargin=10*mm,rightMargin=10*mm,topMargin=10*mm,bottomMargin=10*mm,title='فیش حقوقی').build(story);return b.getvalue()
