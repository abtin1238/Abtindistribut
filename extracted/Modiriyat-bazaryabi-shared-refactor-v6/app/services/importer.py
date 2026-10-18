"""ورود خودکار اطلاعات از خروجی هلو/اکسل/CSV/HTML/XML/JSON.
این ماژول فقط خواندن فایل، تشخیص سرستون/نوع داده، نگاشت ستون‌ها و اعتبارسنجی مقادیر را انجام می‌دهد (بدون وابستگی به DB)."""
from __future__ import annotations
import csv
import io
import json
import re
import zipfile
from datetime import date,datetime
from decimal import Decimal,InvalidOperation
from html.parser import HTMLParser

MAX_ROWS=20000
MAX_COLS=80
HEADER_SCAN=40
_DIG=str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩','01234567890123456789')
_JUNK=dict.fromkeys(map(ord,'\u200c\u200d\u200e\u200f\u202a\u202b\u202c\u0640\ufeff'),None)
_PUNCT=re.compile(r'[\s:;()\[\]{}.,،؛_\-/\\*#"\'«»]+')
class ImportError_(ValueError): pass

def fa_norm(x):
 t=str(x).translate(_DIG).translate(_JUNK).replace('ي','ی').replace('ك','ک').replace('ى','ی').replace('ۀ','ه').replace('\xa0',' ')
 return re.sub(r'\s+',' ',t).strip()
def key_norm(x): return _PUNCT.sub('',fa_norm(x).lower())
def cell_text(v):
 if v is None: return ''
 if isinstance(v,bool): return 'بله' if v else 'خیر'
 if isinstance(v,int): return str(v)
 if isinstance(v,float):
  if v!=v or v in (float('inf'),float('-inf')): return ''
  return str(int(v)) if v.is_integer() and abs(v)<1e15 else format(Decimal(repr(v)),'f')
 if isinstance(v,Decimal): return format(v,'f')
 if isinstance(v,(datetime,date)): return v.date().isoformat() if isinstance(v,datetime) else v.isoformat()
 return fa_norm(v)

# ───────────── خواندن فایل‌ها ─────────────
def _decode(data):
 if data[:2] in (b'\xff\xfe',b'\xfe\xff'): return data.decode('utf-16')
 for enc in ('utf-8-sig','cp1256'):
  try: return data.decode(enc)
  except UnicodeDecodeError: pass
 return data.decode('latin-1')
def _trim(rows):
 out=[]
 for r in rows:
  r=list(r)[:MAX_COLS]
  while r and (r[-1] is None or cell_text(r[-1])==''): r.pop()
  out.append(r)
  if len(out)>MAX_ROWS+HEADER_SCAN: raise ImportError_(f'تعداد سطرهای فایل بیش از {MAX_ROWS} است؛ فایل را تقسیم کنید')
 return out
def _read_xlsx(data):
 try:
  zf=zipfile.ZipFile(io.BytesIO(data))
  if sum(i.file_size for i in zf.infolist())>400*1024*1024: raise ImportError_('فایل اکسل بیش از حد بزرگ است')
  import openpyxl
  wb=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True)
 except ImportError_: raise
 except Exception as e: raise ImportError_('فایل اکسل (xlsx) معتبر نیست یا خراب است') from e
 out={}
 for ws in wb.worksheets:
  if getattr(ws,'sheet_state','visible')!='visible': continue
  rows=_trim(ws.iter_rows(values_only=True))
  if any(rows): out[ws.title]=rows
 wb.close();return out
def _read_xls(data):
 try: import xlrd
 except ImportError as e: raise ImportError_('فایل xls قدیمی است؛ آن را در اکسل با Save As به xlsx یا CSV تبدیل کنید (یا بسته‌ی xlrd را نصب کنید)') from e
 try: wb=xlrd.open_workbook(file_contents=data)
 except Exception as e: raise ImportError_('فایل xls معتبر نیست') from e
 out={}
 for sh in wb.sheets():
  rows=_trim([[ (xlrd.xldate_as_datetime(c.value,wb.datemode) if c.ctype==xlrd.XL_CELL_DATE else c.value) if c.ctype not in (xlrd.XL_CELL_EMPTY,xlrd.XL_CELL_BLANK,xlrd.XL_CELL_ERROR) else None for c in sh.row(i)] for i in range(sh.nrows)])
  if any(rows): out[sh.name]=rows
 return out
class _Tables(HTMLParser):
 def __init__(s): super().__init__(convert_charrefs=True);s.tables=[];s.cur=None;s.row=None;s.cell=None;s.depth=0
 def handle_starttag(s,t,a):
  if t=='table':
   s.depth+=1
   if s.depth==1: s.cur=[]
  elif s.cur is not None and s.depth==1:
   if t=='tr': s.row=[]
   elif t in ('td','th') and s.row is not None: s.cell=[]
   elif t=='br' and s.cell is not None: s.cell.append(' ')
 def handle_endtag(s,t):
  if t=='table':
   if s.depth==1 and s.cur is not None: s.tables.append(s.cur);s.cur=None
   s.depth=max(0,s.depth-1)
  elif s.depth==1:
   if t in ('td','th') and s.cell is not None and s.row is not None: s.row.append(''.join(s.cell));s.cell=None
   elif t=='tr' and s.row is not None and s.cur is not None: s.cur.append(s.row);s.row=None
 def handle_data(s,d):
  if s.cell is not None: s.cell.append(d)
def _read_html(text):
 p=_Tables()
 try: p.feed(text);p.close()
 except Exception as e: raise ImportError_('جدول HTML قابل خواندن نیست') from e
 return {f'جدول {i+1}':_trim(t) for i,t in enumerate(p.tables) if len(t)>=2}
def _read_xml2003(text):
 if '<!DOCTYPE' in text.upper() or '<!ENTITY' in text.upper(): raise ImportError_('فایل XML مجاز نیست')
 import xml.etree.ElementTree as ET
 try: root=ET.fromstring(text.encode('utf-8'))
 except ET.ParseError as e: raise ImportError_('فایل XML اکسل معتبر نیست') from e
 loc=lambda e:e.tag.rsplit('}',1)[-1];out={}
 for ws in root.iter():
  if loc(ws)!='Worksheet': continue
  rows=[]
  for row in ws.iter():
   if loc(row)!='Row': continue
   r=[]
   for c in row:
    if loc(c)!='Cell': continue
    idx=[v for k,v in c.attrib.items() if k.endswith('Index')]
    if idx: r+= [None]*(int(idx[0])-1-len(r))
    d=next((x for x in c if loc(x)=='Data'),None);r.append(''.join(d.itertext()) if d is not None else None)
   rows.append(r)
  rows=_trim(rows)
  if any(rows): out[next((v for k,v in ws.attrib.items() if k.endswith('Name')),f'Sheet{len(out)+1}')]=rows
 return out
def _read_json(text):
 try: d=json.loads(text)
 except ValueError as e: raise ImportError_('فایل JSON معتبر نیست') from e
 if isinstance(d,dict): d=next((v for v in d.values() if isinstance(v,list) and v and isinstance(v[0],dict)),None)
 if not isinstance(d,list) or not d or not all(isinstance(x,dict) for x in d): raise ImportError_('JSON باید آرایه‌ای از اشیاء باشد')
 cols=[];[cols.append(k) for x in d for k in x if k not in cols]
 return {'JSON':_trim([cols]+[[x.get(k) if not isinstance(x.get(k),(dict,list)) else json.dumps(x.get(k),ensure_ascii=False) for k in cols] for x in d])}
def _read_csv(text):
 sample=text[:20000];lines=[l for l in sample.splitlines() if l.strip()][:30]
 best=max(',;\t|',key=lambda d:sum(l.count(d) for l in lines))
 if not any(best in l for l in lines): best=','
 rows=[r for r in csv.reader(io.StringIO(text),delimiter=best)]
 return {'CSV':_trim(rows)}
def _is_dbf(d):
 if len(d)<33 or d[0] not in (0x02,0x03,0x04,0x05,0x30,0x31,0x32,0x83,0x8b,0xf5,0xfb): return False
 hl=int.from_bytes(d[8:10],'little');rl=int.from_bytes(d[10:12],'little')
 return 33<=hl<=65535 and rl>=2 and (hl-1-32)%32==0 and d[hl-1:hl]==b'\r' and 1<=(hl-33)//32<=255
def _dbf_text(b):
 b=b.split(b'\x00')[0].rstrip(b' ')
 for enc in ('utf-8','cp1256'):
  try: return fa_norm(b.decode(enc))
  except UnicodeDecodeError: pass
 return fa_norm(b.decode('latin-1'))
def _read_dbf(d):
 n=int.from_bytes(d[4:8],'little');hl=int.from_bytes(d[8:10],'little');rl=int.from_bytes(d[10:12],'little');cols=[]
 for o in range(32,hl-1,32):
  f=d[o:o+32];cols.append((f[:11].split(b'\x00')[0].decode('latin-1').strip(),chr(f[11]),f[16],f[17]))
 off=0;sl=[]
 for c in cols: sl.append((off+1,off+1+c[2]));off+=c[2]
 rows=[[c[0] for c in cols]]
 if n>MAX_ROWS+HEADER_SCAN: raise ImportError_(f'تعداد سطرهای فایل بیش از {MAX_ROWS} است؛ فایل را تقسیم کنید')
 for i in range(n):
  rec=d[hl+i*rl:hl+(i+1)*rl]
  if len(rec)<rl: break
  if rec[:1]==b'*': continue  # حذف‌شده
  r=[]
  for c,(a,b) in zip(cols,sl):
   raw=rec[a:b];t=c[1]
   if t in 'NF':
    s=raw.decode('latin-1').strip();r.append(None if not s or set(s)<=set('*?') else s)
   elif t=='L': r.append({'T':True,'Y':True,'F':False,'N':False}.get(raw.decode('latin-1').strip().upper()[:1]))
   elif t=='D':
    s=raw.decode('latin-1').strip();r.append(f'{s[:4]}-{s[4:6]}-{s[6:8]}' if len(s)==8 and s.isdigit() else None)
   elif t in 'MBGP+': r.append(None)
   else: r.append(_dbf_text(raw))
  rows.append(r)
 return {'DBF':_trim(rows)}
def _read_ods(data):
 import xml.etree.ElementTree as ET
 try:
  zf=zipfile.ZipFile(io.BytesIO(data));raw=zf.read('content.xml')
 except Exception as e: raise ImportError_('فایل ODS معتبر نیست') from e
 if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper(): raise ImportError_('فایل ODS مجاز نیست')
 try: root=ET.fromstring(raw)
 except ET.ParseError as e: raise ImportError_('فایل ODS معتبر نیست') from e
 loc=lambda e:e.tag.rsplit('}',1)[-1];at=lambda e,n:next((v for k,v in e.attrib.items() if k.endswith('}'+n)),None);out={}
 for tb in root.iter():
  if loc(tb)!='table': continue
  rows=[]
  for row in tb.iter():
   if loc(row)!='table-row': continue
   r=[]
   for c in row:
    if loc(c) not in ('table-cell','covered-table-cell'): continue
    v=at(c,'value');v=''.join(''.join(p.itertext()) for p in c if loc(p)=='p') if v is None else v
    rep=min(int(at(c,'number-columns-repeated') or 1),MAX_COLS);r+=[v if v!='' else None]*rep
   rr=min(int(at(row,'number-rows-repeated') or 1),1 if not any(x is not None for x in r) else 50)
   rows+=[r]*rr
  rows=_trim(rows)
  if any(rows): out[at(tb,'name') or f'Sheet{len(out)+1}']=rows
 return out
def load_tables(data,filename=''):
 """بر اساس محتوا (نه پسوند) نوع فایل را تشخیص می‌دهد؛ خروجی: {نام برگه: [[سلول...]...]}"""
 if not data.strip(): raise ImportError_('فایل خالی است')
 if data[:4]==b'PK\x03\x04': out=_read_ods(data) if b'opendocument.spreadsheet' in data[:600] else _read_xlsx(data)
 elif _is_dbf(data): out=_read_dbf(data)
 elif data[:8]==b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1': out=_read_xls(data)
 else:
  text=_decode(data);head=text.lstrip('\ufeff \t\r\n')[:400].lower()
  if 'urn:schemas-microsoft-com:office:spreadsheet' in text[:2000]: out=_read_xml2003(text)
  elif head.startswith('<'): out=_read_html(text)
  elif head[:1] in '{[': out=_read_json(text)
  else: out=_read_csv(text)
 if not out: raise ImportError_('داده‌ای در فایل پیدا نشد')
 return out

# ───────────── تعریف موجودیت‌ها و مترادف سرستون‌ها ─────────────
def _k(*xs): return [key_norm(x) for x in xs]
GENERIC=set(_k('کد','code','نام','name','عنوان','شرح','description','فی','قیمت','price','مانده','تعداد','مقدار','وضعیت','گروه','انبار'))
_PHONE=_k('تلفن','موبایل','همراه','تلفن همراه','شماره تماس','شماره موبایل','شماره تلفن','تلفن تماس','تلفن ثابت','تلفن 1','تلفن1','phone','mobile','tel','tel1','cell')
_ADDR=_k('آدرس','ادرس','نشانی','نشانی کامل','آدرس کامل','address')
_STOCK_Q=_k('موجودی','موجودی فعلی','موجودی انبار','موجودی کل','تعداد موجود','مقدار موجود','موجودی نهایی','تعداد','مقدار','qty','quantity','stock','onhand','on hand','exist','mojoodi')
_WH=_k('انبار','نام انبار','محل نگهداری','انبار مقصد','warehouse','store')
_BAL=_k('مانده','مانده حساب','مانده نهایی','مانده کل','مانده (ریال)','مانده (تومان)','مانده ریالی','مانده نقدی','balance')
_DEB=_k('بدهکار','گردش بدهکار','مانده بدهکار','جمع بدهکار','بد','debit')
_CRE=_k('بستانکار','گردش بستانکار','مانده بستانکار','جمع بستانکار','بس','credit')
_BT=_k('وضعیت','تشخیص','ماهیت','بد/بس','بد بس','نوع مانده','وضعیت حساب','بدهکار/بستانکار','تشخیص مانده')
ENTITIES={
 'products':{'label':'کالا و قیمت (با موجودی اختیاری)','key':'sku','fields':{
  'sku':('کد کالا',True,_k('کد کالا','کد محصول','شماره کالا','کد کالا/خدمات','کد کالا و خدمات','کد جنس','کد','code','sku','item code','product code','a_code','acode2')),
  'name':('نام کالا',True,_k('نام کالا','شرح کالا','عنوان کالا','نام محصول','شرح کالا/خدمات','نام جنس','شرح جنس','نام','عنوان','شرح','name','item name','product name','description','a_name')),
  'unit':('واحد',False,_k('واحد','واحد اصلی','واحد شمارش','واحد کالا','واحد سنجش','unit','uom')),
  'base_price':('قیمت فروش',False,_k('قیمت فروش','فی فروش','قیمت واحد','قیمت مصرف کننده','قیمت','فی','مبلغ واحد','قیمت فروش 1','price','sale price','unit price','sel_price','sell_price')),
  'min_stock':('حداقل موجودی',False,_k('حداقل موجودی','نقطه سفارش','حداقل','سفارش مجدد','min stock','reorder point')),
  'quantity':('موجودی',False,_STOCK_Q),
  'warehouse':('انبار',False,_WH),
  'brand':('برند',False,_k('برند','نشان تجاری','سازنده','تولید کننده','brand')),
  'category':('گروه کالا',False,_k('گروه کالا','دسته بندی','گروه اصلی','نام گروه','گروه','دسته','category'))}},
 'stock':{'label':'موجودی انبار (کالاهای موجود)','key':'sku','fields':{
  'sku':('کد کالا',True,_k('کد کالا','کد محصول','شماره کالا','کد کالا/خدمات','کد جنس','کد','code','sku','item code','a_code')),
  'name':('نام کالا',False,_k('نام کالا','شرح کالا','عنوان کالا','نام محصول','نام جنس','نام','عنوان','شرح','name','item name','description','a_name')),
  'quantity':('موجودی',True,_STOCK_Q),
  'warehouse':('انبار',False,_WH)}},
 'customers':{'label':'مشتریان و مانده حساب','key':'code','fields':{
  'code':('کد مشتری',False,_k('کد مشتری','کد حساب','کد طرف حساب','کد تفصیلی','کد معین','شماره حساب','کد','code','customer code','c_code')),
  'name':('نام مشتری',True,_k('نام مشتری','نام طرف حساب','نام حساب','نام و نام خانوادگی','عنوان حساب','شرح حساب','نام شرکت','مشتری','طرف حساب','نام','name','customer name','c_name')),
  'phone':('تلفن',False,_PHONE),'address':('آدرس',False,_ADDR),
  'region':('منطقه/شهر',False,_k('منطقه','شهر','ناحیه','استان','محله','مسیر','خط سیر','region','city')),
  'customer_group':('گروه مشتری',False,_k('گروه مشتری','گروه طرف حساب','دسته مشتری','سطح مشتری','گروه','customer group','group')),
  'credit_limit':('سقف اعتبار',False,_k('سقف اعتبار','سقف بدهی','حداکثر اعتبار','حد مجاز اعتبار','اعتبار','credit limit')),
  'balance':('مانده حساب',False,_BAL),'debit':('بدهکار',False,_DEB),'credit':('بستانکار',False,_CRE),'balance_type':('بد/بس مانده',False,_BT)}},
 'suppliers':{'label':'تأمین‌کنندگان و مانده حساب','key':'name','fields':{
  'name':('نام تأمین‌کننده',True,_k('نام تامین کننده','نام فروشنده','تامین کننده','فروشنده','نام طرف حساب','نام حساب','نام شرکت','نام','name','supplier')),
  'phone':('تلفن',False,_PHONE),
  'balance':('مانده حساب',False,_BAL),'debit':('بدهکار',False,_DEB),'credit':('بستانکار',False,_CRE),'balance_type':('بد/بس مانده',False,_BT)}}}
NEG={'base_price':_k('خرید','تمام شده','میانگین','cost','buy','حداقل','حداکثر'),'quantity':_k('حداقل','حداکثر','نقطه','قیمت','فی','مبلغ'),'name':_k('کد','شماره'),'region':_k('کد')}
def _score(field,h,syns):
 if not h: return 0
 if any(n in h for n in NEG.get(field,())): return 0
 best=0
 for s in syns:
  if h==s: best=max(best,1.5 if s in GENERIC else 3)
  elif len(s)>=4 and s in h and s not in GENERIC: best=max(best,2)
 return best
def map_headers(entity,headers):
 """{ستون: فیلد} با انتخاب حریصانه‌ی بهترین امتیاز؛ هر فیلد/ستون حداکثر یک‌بار."""
 fields=ENTITIES[entity]['fields'];hs=[key_norm(h) for h in headers];cand=[]
 for i,h in enumerate(hs):
  for f,(_,_,syn) in fields.items():
   sc=_score(f,h,syn)
   if sc>0: cand.append((-sc,i,f))
 cand.sort();used_c=set();used_f=set();out={}
 for nsc,i,f in cand:
  if i in used_c or f in used_f: continue
  out[i]=(f,-nsc);used_c.add(i);used_f.add(f)
 return out
def detect(rows):
 """(ردیف سرستون، موجودیت، امتیاز)؛ ردیف عنوان/لوگوی بالای گزارش‌های هلو نادیده گرفته می‌شود."""
 best=(0,None,0.0)
 for ri,r in enumerate(rows[:HEADER_SCAN]):
  if sum(1 for c in r if cell_text(c))<2: continue
  hs=[cell_text(c) for c in r]
  for e in ENTITIES:
   m=map_headers(e,hs);keyf=ENTITIES[e]['key'];req=[f for f,(_,rq,_) in ENTITIES[e]['fields'].items() if rq]
   got={f for f,_ in m.values()}
   if not all(f in got for f in req) and keyf not in got and 'name' not in got: continue
   sc=sum(s for _,s in m.values())+(1 if all(f in got for f in req) else 0)
   if sc>best[2]: best=(ri,e,sc)
 return best

# ───────────── تبدیل مقادیر ─────────────
_CUR=re.compile(r'(ریال|تومان|rials?|irr|toman)$')
def parse_num(v):
 """Decimal یا None (خالی)؛ ValueError برای مقدار نامعتبر. کاما/٬ جداکننده‌ی هزارگان است؛ (123) و 123- و −123 منفی."""
 if v is None: return None
 if isinstance(v,bool): raise ValueError('عدد نامعتبر')
 if isinstance(v,(int,float,Decimal)):
  d=Decimal(str(v))
  if not d.is_finite(): raise ValueError('عدد نامعتبر')
  return d
 t=fa_norm(v).lower()
 if t in ('','-','—','–'): return None
 neg=False;t=t.replace('٬','').replace(',','').replace(' ','').replace('٫','.').replace('−','-').replace('–','-')
 t=_CUR.sub('',t)
 if t.startswith('(') and t.endswith(')'): neg=True;t=t[1:-1]
 if t.endswith('-'): neg=not neg;t=t[:-1]
 if t.startswith('-'): neg=not neg;t=t[1:]
 t=t.lstrip('+')
 if not re.fullmatch(r'\d+(\.\d+)?|\.\d+',t): raise ValueError('عدد نامعتبر: «%s»'%fa_norm(v)[:30])
 d=Decimal(t);return -d if neg else d
def norm_phone(v):
 t=cell_text(v)
 if not t: return ''
 t=re.split(r'[/،,;؛]|\s-\s',t)[0].strip();d=re.sub(r'[\s\-()]','',t)
 if d.startswith('+98'): d='0'+d[3:]
 elif d.startswith('0098'): d='0'+d[4:]
 elif d.startswith('98') and len(d)==12: d='0'+d[2:]
 elif len(d)==10 and d.startswith('9'): d='0'+d
 return d if re.fullmatch(r'\d{6,15}',d) else t
LIMITS={'sku':80,'name':200,'unit':30,'phone':50,'code':60,'region':120,'customer_group':100,'brand':150,'category':150,'warehouse':200,'address':2000}
NUMF={'base_price':'money','min_stock':'qty','quantity':'qty','credit_limit':'money','balance':'signed','debit':'money','credit':'money'}
def _bt(v):
 t=key_norm(v)
 if not t: return None
 if t.startswith(('بد','debit','dr','d')): return 'dr'
 if t.startswith(('بس','credit','cr','c')): return 'cr'
 raise ValueError('مقدار «بد/بس» نامعتبر')
def coerce_rows(entity,rows,header_row,mapping):
 """mapping: {ستون(int): فیلد}. خروجی: لیست dict {line,values,errors,skip} ؛ line شماره‌ی ردیف در فایل (۱-مبنا)."""
 spec=ENTITIES[entity];inv={f:i for i,f in mapping.items() if f in spec['fields']};hdr_keys={key_norm(cell_text(c)) for c in rows[header_row] if cell_text(c)}
 out=[];seen={}
 for ri in range(header_row+1,len(rows)):
  r=rows[ri];cells=[cell_text(c) for c in r]
  if not any(cells): continue
  ne=[c for c in cells if c]
  rk={key_norm(c) for c in ne}
  if len(rk&hdr_keys)>=max(2,len(ne)//2+1): continue  # سرستون تکراری هر صفحه‌ی گزارش
  g=lambda f:(r[inv[f]] if f in inv and inv[f]<len(r) else None)
  vals={};errs=[]
  for f in inv:
   raw=g(f);t=cell_text(raw)
   if f in NUMF:
    try: d=parse_num(raw);vals[f]=d
    except ValueError as e: errs.append(f'{spec["fields"][f][0]}: {e}');continue
    if d is not None and f in ('base_price','min_stock','quantity','credit_limit','debit','credit') and d<0: errs.append(f'{spec["fields"][f][0]} نمی‌تواند منفی باشد')
   elif f=='phone': vals[f]=norm_phone(raw)
   elif f=='balance_type':
    try: vals[f]=_bt(raw)
    except ValueError as e: errs.append(str(e))
   else:
    if t.startswith("'"): t=t[1:].strip()
    vals[f]=t
    if len(t)>LIMITS.get(f,200): errs.append(f'{spec["fields"][f][0]} بیش از {LIMITS.get(f,200)} نویسه است')
  keyf=spec['key'];kv=vals.get(keyf) or ''
  if not kv and keyf=='code': kv=vals.get('name') or ''
  empty_key=not kv and not vals.get('name') and not vals.get('sku')
  if empty_key or (len(ne)<2 and not kv): out.append({'line':ri+1,'values':vals,'errors':[],'skip':'عنوان گروه/جمع/ردیف خالی'});continue
  if re.match(r'^(جمع|مجموع|total|sum)\b',fa_norm(next(iter(ne))).lower()) and not vals.get('sku'): out.append({'line':ri+1,'values':vals,'errors':[],'skip':'ردیف جمع'});continue
  for f,(lab,rq,_) in spec['fields'].items():
   if rq and f in inv and not vals.get(f) and vals.get(f)!=Decimal(0): errs.append(f'{lab} خالی است')
  if 'sku' in spec['fields'] and 'sku' not in inv and entity in ('products','stock'): errs.append('ستون «کد کالا» نگاشت نشده است')
  if entity in ('customers','suppliers'):
   d,c,b,bt=vals.get('debit'),vals.get('credit'),vals.get('balance'),vals.get('balance_type')
   amt=None
   if b is not None:
    amt=b
    if bt:  # بد/بس صریح: مبلغ مثبت = بدهکار طرف مشتری / بستانکار طرف تأمین‌کننده
     amt=abs(b)*(1 if (bt=='dr')==(entity=='customers') else -1)
   elif d is not None or c is not None:
    amt=(d or Decimal(0))-(c or Decimal(0)) if entity=='customers' else (c or Decimal(0))-(d or Decimal(0))
   vals['_balance']=amt
  dk=(key_norm(kv),key_norm(vals.get('warehouse') or '')) if entity=='stock' else key_norm(kv)
  if kv and dk in seen: errs.append(f'تکراری در فایل (ردیف {seen[dk]})')
  elif kv: seen[dk]=ri+1
  out.append({'line':ri+1,'values':vals,'errors':errs,'skip':''})
 return out
